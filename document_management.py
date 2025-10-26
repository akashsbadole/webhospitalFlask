#!/usr/bin/env python3
"""
Document Management System
Handles document upload, versioning, digital signatures, templates, OCR, and expiry tracking
"""

import os
import hashlib
import mimetypes
from datetime import datetime, timedelta
from pathlib import Path
import json
import logging
from flask import current_app, request, session
from flask_login import current_user
from werkzeug.utils import secure_filename
from PIL import Image
import PyPDF2
from docx import Document as DocxDocument
import pytesseract
import zipfile
from cryptography.fernet import Fernet
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.asymmetric import rsa, padding
from cryptography.hazmat.primitives import serialization

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class DocumentManagement:
    def __init__(self, db, app):
        self.db = db
        self.app = app
        self.upload_folder = Path(app.config.get('UPLOAD_FOLDER', 'uploads'))
        self.upload_folder.mkdir(exist_ok=True)
        
        # Create subdirectories
        (self.upload_folder / 'documents').mkdir(exist_ok=True)
        (self.upload_folder / 'templates').mkdir(exist_ok=True)
        (self.upload_folder / 'signatures').mkdir(exist_ok=True)
        (self.upload_folder / 'temp').mkdir(exist_ok=True)
        
        self.setup_document_tables()
        
        # Configuration
        self.config = {
            'max_file_size': 16 * 1024 * 1024,  # 16MB
            'allowed_extensions': {
                'pdf', 'doc', 'docx', 'jpg', 'jpeg', 'png', 'gif', 'txt', 'rtf'
            },
            'virus_scan_enabled': False,  # Configure in production
            'ocr_enabled': True,
            'encryption_enabled': True,
            'digital_signature_enabled': True,
            'version_retention_days': 365,
            'auto_expiry_check': True
        }
        
        # Generate encryption key if not exists
        self.encryption_key = self._get_or_create_encryption_key()
    
    def setup_document_tables(self):
        """Create document management tables"""
        
        class Document(self.db.Model):
            __tablename__ = 'documents'
            
            id = self.db.Column(self.db.Integer, primary_key=True)
            document_id = self.db.Column(self.db.String(50), unique=True, nullable=False)
            title = self.db.Column(self.db.String(200), nullable=False)
            description = self.db.Column(self.db.Text)
            category = self.db.Column(self.db.String(50), nullable=False)  # consent, insurance, medical, etc.
            document_type = self.db.Column(self.db.String(50))  # discharge_summary, consent_form, etc.
            
            # File information
            filename = self.db.Column(self.db.String(255), nullable=False)
            original_filename = self.db.Column(self.db.String(255), nullable=False)
            file_path = self.db.Column(self.db.String(500), nullable=False)
            file_size = self.db.Column(self.db.Integer)
            mime_type = self.db.Column(self.db.String(100))
            file_hash = self.db.Column(self.db.String(64))  # SHA-256 hash
            
            # Relationships
            patient_id = self.db.Column(self.db.Integer, self.db.ForeignKey('web_patient.id'))
            staff_id = self.db.Column(self.db.Integer, self.db.ForeignKey('web_user.id'))
            uploaded_by = self.db.Column(self.db.Integer, self.db.ForeignKey('web_user.id'), nullable=False)
            
            # Versioning
            version = self.db.Column(self.db.Integer, default=1)
            parent_document_id = self.db.Column(self.db.Integer, self.db.ForeignKey('documents.id'))
            is_current_version = self.db.Column(self.db.Boolean, default=True)
            
            # Status and metadata
            status = self.db.Column(self.db.String(20), default='active')  # active, archived, deleted
            is_encrypted = self.db.Column(self.db.Boolean, default=False)
            is_signed = self.db.Column(self.db.Boolean, default=False)
            
            # Expiry and compliance
            expiry_date = self.db.Column(self.db.Date)
            retention_period_years = self.db.Column(self.db.Integer)
            compliance_tags = self.db.Column(self.db.Text)  # JSON array
            
            # OCR and search
            ocr_text = self.db.Column(self.db.Text)
            searchable_content = self.db.Column(self.db.Text)
            
            # Audit fields
            created_at = self.db.Column(self.db.DateTime, default=datetime.utcnow)
            updated_at = self.db.Column(self.db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
            accessed_at = self.db.Column(self.db.DateTime)
            access_count = self.db.Column(self.db.Integer, default=0)
            
            # Relationships
            patient = self.db.relationship('WebPatient', foreign_keys=[patient_id], backref='documents')
            staff = self.db.relationship('WebUser', foreign_keys=[staff_id])
            uploader = self.db.relationship('WebUser', foreign_keys=[uploaded_by])
            parent_document = self.db.relationship('Document', remote_side=[id], backref='versions')
        
        class DocumentTemplate(self.db.Model):
            __tablename__ = 'document_templates'
            
            id = self.db.Column(self.db.Integer, primary_key=True)
            name = self.db.Column(self.db.String(100), nullable=False)
            category = self.db.Column(self.db.String(50), nullable=False)
            template_type = self.db.Column(self.db.String(50))  # discharge, consent, certificate
            
            # Template content
            content = self.db.Column(self.db.Text)  # HTML/Rich text content
            variables = self.db.Column(self.db.Text)  # JSON array of template variables
            
            # File template (if applicable)
            template_file_path = self.db.Column(self.db.String(500))
            
            # Metadata
            description = self.db.Column(self.db.Text)
            is_active = self.db.Column(self.db.Boolean, default=True)
            requires_signature = self.db.Column(self.db.Boolean, default=False)
            
            # Audit
            created_by = self.db.Column(self.db.Integer, self.db.ForeignKey('web_user.id'))
            created_at = self.db.Column(self.db.DateTime, default=datetime.utcnow)
            updated_at = self.db.Column(self.db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
            
            creator = self.db.relationship('WebUser', backref='created_templates')
        
        class DocumentSignature(self.db.Model):
            __tablename__ = 'document_signatures'
            
            id = self.db.Column(self.db.Integer, primary_key=True)
            document_id = self.db.Column(self.db.Integer, self.db.ForeignKey('documents.id'), nullable=False)
            signer_id = self.db.Column(self.db.Integer, self.db.ForeignKey('web_user.id'), nullable=False)
            
            # Signature data
            signature_type = self.db.Column(self.db.String(20), default='digital')  # digital, electronic, wet
            signature_data = self.db.Column(self.db.Text)  # Base64 encoded signature or certificate
            signature_hash = self.db.Column(self.db.String(128))
            
            # Verification
            is_verified = self.db.Column(self.db.Boolean, default=False)
            verification_method = self.db.Column(self.db.String(50))
            
            # Metadata
            signed_at = self.db.Column(self.db.DateTime, default=datetime.utcnow)
            ip_address = self.db.Column(self.db.String(45))
            user_agent = self.db.Column(self.db.Text)
            
            # Relationships
            document = self.db.relationship('Document', backref='signatures')
            signer = self.db.relationship('WebUser', backref='signed_documents')
        
        class DocumentAccess(self.db.Model):
            __tablename__ = 'document_access'
            
            id = self.db.Column(self.db.Integer, primary_key=True)
            document_id = self.db.Column(self.db.Integer, self.db.ForeignKey('documents.id'), nullable=False)
            user_id = self.db.Column(self.db.Integer, self.db.ForeignKey('web_user.id'), nullable=False)
            
            # Access details
            access_type = self.db.Column(self.db.String(20), nullable=False)  # view, download, edit
            accessed_at = self.db.Column(self.db.DateTime, default=datetime.utcnow)
            ip_address = self.db.Column(self.db.String(45))
            user_agent = self.db.Column(self.db.Text)
            
            # Relationships
            document = self.db.relationship('Document', backref='access_logs')
            user = self.db.relationship('WebUser', backref='document_accesses')
        
        # Store model classes
        self.Document = Document
        self.DocumentTemplate = DocumentTemplate
        self.DocumentSignature = DocumentSignature
        self.DocumentAccess = DocumentAccess
        
        # Create tables
        self.db.create_all()
        self._create_default_templates()
    
    def _create_default_templates(self):
        """Create default document templates"""
        default_templates = [
            {
                'name': 'Discharge Summary',
                'category': 'medical',
                'template_type': 'discharge',
                'content': '''
                <h2>DISCHARGE SUMMARY</h2>
                <p><strong>Patient Name:</strong> {{patient_name}}</p>
                <p><strong>Patient ID:</strong> {{patient_id}}</p>
                <p><strong>Admission Date:</strong> {{admission_date}}</p>
                <p><strong>Discharge Date:</strong> {{discharge_date}}</p>
                <p><strong>Attending Physician:</strong> {{doctor_name}}</p>
                
                <h3>Diagnosis</h3>
                <p>{{diagnosis}}</p>
                
                <h3>Treatment Provided</h3>
                <p>{{treatment}}</p>
                
                <h3>Discharge Instructions</h3>
                <p>{{instructions}}</p>
                
                <h3>Follow-up</h3>
                <p>{{followup}}</p>
                
                <p><strong>Physician Signature:</strong> ________________</p>
                <p><strong>Date:</strong> {{current_date}}</p>
                ''',
                'variables': json.dumps([
                    'patient_name', 'patient_id', 'admission_date', 'discharge_date',
                    'doctor_name', 'diagnosis', 'treatment', 'instructions', 'followup', 'current_date'
                ]),
                'requires_signature': True
            },
            {
                'name': 'Consent Form',
                'category': 'consent',
                'template_type': 'consent',
                'content': '''
                <h2>INFORMED CONSENT FOR MEDICAL TREATMENT</h2>
                
                <p>I, {{patient_name}}, hereby give my informed consent for the medical treatment/procedure described below:</p>
                
                <p><strong>Procedure/Treatment:</strong> {{procedure_name}}</p>
                <p><strong>Doctor:</strong> {{doctor_name}}</p>
                <p><strong>Date:</strong> {{procedure_date}}</p>
                
                <h3>Nature of the Procedure</h3>
                <p>{{procedure_description}}</p>
                
                <h3>Risks and Complications</h3>
                <p>{{risks}}</p>
                
                <h3>Alternative Treatments</h3>
                <p>{{alternatives}}</p>
                
                <p>I acknowledge that I have read and understood this consent form, and all my questions have been answered to my satisfaction.</p>
                
                <p><strong>Patient Signature:</strong> ________________ <strong>Date:</strong> ________</p>
                <p><strong>Witness Signature:</strong> ________________ <strong>Date:</strong> ________</p>
                <p><strong>Doctor Signature:</strong> ________________ <strong>Date:</strong> ________</p>
                ''',
                'variables': json.dumps([
                    'patient_name', 'procedure_name', 'doctor_name', 'procedure_date',
                    'procedure_description', 'risks', 'alternatives'
                ]),
                'requires_signature': True
            },
            {
                'name': 'Medical Certificate',
                'category': 'certificate',
                'template_type': 'certificate',
                'content': '''
                <h2>MEDICAL CERTIFICATE</h2>
                
                <p>This is to certify that {{patient_name}} (Age: {{age}}) was under my medical care from {{start_date}} to {{end_date}}.</p>
                
                <p><strong>Diagnosis:</strong> {{diagnosis}}</p>
                
                <p><strong>Recommendation:</strong> {{recommendation}}</p>
                
                <p>This certificate is issued at the request of the patient for {{purpose}}.</p>
                
                <p><strong>Doctor:</strong> {{doctor_name}}</p>
                <p><strong>Registration No:</strong> {{doctor_license}}</p>
                <p><strong>Date:</strong> {{issue_date}}</p>
                
                <p><strong>Signature:</strong> ________________</p>
                ''',
                'variables': json.dumps([
                    'patient_name', 'age', 'start_date', 'end_date', 'diagnosis',
                    'recommendation', 'purpose', 'doctor_name', 'doctor_license', 'issue_date'
                ]),
                'requires_signature': True
            }
        ]
        
        for template_data in default_templates:
            existing = self.DocumentTemplate.query.filter_by(name=template_data['name']).first()
            if not existing:
                template = self.DocumentTemplate(**template_data)
                self.db.session.add(template)
        
        try:
            self.db.session.commit()
        except Exception as e:
            logger.error(f"Failed to create default templates: {str(e)}")
            self.db.session.rollback()
    
    def upload_document(self, file, title, category, patient_id=None, staff_id=None, 
                       description=None, document_type=None, expiry_date=None):
        """Upload and process a document"""
        try:
            if not file or not file.filename:
                raise ValueError("No file provided")
            
            # Validate file
            if not self._is_allowed_file(file.filename):
                raise ValueError("File type not allowed")
            
            if file.content_length and file.content_length > self.config['max_file_size']:
                raise ValueError("File size exceeds limit")
            
            # Generate unique filename
            original_filename = secure_filename(file.filename)
            file_extension = Path(original_filename).suffix
            document_id = self._generate_document_id()
            filename = f"{document_id}{file_extension}"
            
            # Create file path
            file_path = self.upload_folder / 'documents' / filename
            
            # Save file
            file.save(str(file_path))
            
            # Calculate file hash
            file_hash = self._calculate_file_hash(file_path)
            
            # Get file info
            file_size = file_path.stat().st_size
            mime_type = mimetypes.guess_type(str(file_path))[0]
            
            # Virus scan if enabled
            if self.config['virus_scan_enabled']:
                if not self._scan_for_virus(file_path):
                    file_path.unlink()  # Delete infected file
                    raise ValueError("File failed virus scan")
            
            # Encrypt file if enabled
            encrypted_path = None
            if self.config['encryption_enabled']:
                encrypted_path = self._encrypt_file(file_path)
                if encrypted_path:
                    file_path.unlink()  # Remove original
                    file_path = encrypted_path
            
            # Perform OCR if enabled and applicable
            ocr_text = None
            if self.config['ocr_enabled'] and mime_type in ['image/jpeg', 'image/png', 'application/pdf']:
                ocr_text = self._extract_text_ocr(file_path)
            
            # Create document record
            document = self.Document(
                document_id=document_id,
                title=title,
                description=description,
                category=category,
                document_type=document_type,
                filename=filename,
                original_filename=original_filename,
                file_path=str(file_path),
                file_size=file_size,
                mime_type=mime_type,
                file_hash=file_hash,
                patient_id=patient_id,
                staff_id=staff_id,
                uploaded_by=current_user.id if current_user.is_authenticated else None,
                is_encrypted=self.config['encryption_enabled'],
                expiry_date=expiry_date,
                ocr_text=ocr_text,
                searchable_content=self._create_searchable_content(title, description, ocr_text)
            )
            
            self.db.session.add(document)
            self.db.session.commit()
            
            logger.info(f"Document uploaded successfully: {document_id}")
            return document
            
        except Exception as e:
            logger.error(f"Failed to upload document: {str(e)}")
            raise
    
    def create_document_from_template(self, template_id, data, patient_id=None, 
                                    title=None, requires_signature=False):
        """Create document from template"""
        try:
            template = self.db.session.get(self.DocumentTemplate, template_id)
            if not template:
                raise ValueError("Template not found")
            
            # Render template content
            content = self._render_template_content(template.content, data)
            
            # Generate document title
            if not title:
                title = f"{template.name} - {data.get('patient_name', 'Unknown')}"
            
            # Create HTML file
            document_id = self._generate_document_id()
            filename = f"{document_id}.html"
            file_path = self.upload_folder / 'documents' / filename
            
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(f"""
                <!DOCTYPE html>
                <html>
                <head>
                    <meta charset="UTF-8">
                    <title>{title}</title>
                    <style>
                        body {{ font-family: Arial, sans-serif; margin: 40px; }}
                        h2, h3 {{ color: #333; }}
                        p {{ margin: 10px 0; }}
                    </style>
                </head>
                <body>
                    {content}
                </body>
                </html>
                """)
            
            # Calculate file info
            file_size = file_path.stat().st_size
            file_hash = self._calculate_file_hash(file_path)
            
            # Create document record
            document = self.Document(
                document_id=document_id,
                title=title,
                category=template.category,
                document_type=template.template_type,
                filename=filename,
                original_filename=f"{title}.html",
                file_path=str(file_path),
                file_size=file_size,
                mime_type='text/html',
                file_hash=file_hash,
                patient_id=patient_id,
                uploaded_by=current_user.id if current_user.is_authenticated else None,
                searchable_content=self._create_searchable_content(title, content)
            )
            
            self.db.session.add(document)
            self.db.session.commit()
            
            # If signature required, mark document
            if requires_signature or template.requires_signature:
                document.is_signed = False
            
            logger.info(f"Document created from template: {document_id}")
            return document
            
        except Exception as e:
            logger.error(f"Failed to create document from template: {str(e)}")
            raise
    
    def add_digital_signature(self, document_id, signature_data, signature_type='digital'):
        """Add digital signature to document"""
        try:
            document = self.Document.query.filter_by(document_id=document_id).first()
            if not document:
                raise ValueError("Document not found")
            
            # Create signature hash
            signature_hash = hashlib.sha256(signature_data.encode()).hexdigest()
            
            # Create signature record
            signature = self.DocumentSignature(
                document_id=document.id,
                signer_id=current_user.id if current_user.is_authenticated else None,
                signature_type=signature_type,
                signature_data=signature_data,
                signature_hash=signature_hash,
                ip_address=self._get_client_ip(),
                user_agent=request.headers.get('User-Agent', '') if request else ''
            )
            
            self.db.session.add(signature)
            
            # Update document status
            document.is_signed = True
            
            self.db.session.commit()
            
            logger.info(f"Digital signature added to document: {document_id}")
            return signature
            
        except Exception as e:
            logger.error(f"Failed to add digital signature: {str(e)}")
            raise
    
    def create_new_version(self, document_id, file, description=None):
        """Create new version of existing document"""
        try:
            original_doc = self.Document.query.filter_by(document_id=document_id).first()
            if not original_doc:
                raise ValueError("Original document not found")
            
            # Mark current version as not current
            original_doc.is_current_version = False
            
            # Upload new version
            new_document = self.upload_document(
                file=file,
                title=original_doc.title,
                category=original_doc.category,
                patient_id=original_doc.patient_id,
                staff_id=original_doc.staff_id,
                description=description or f"Version {original_doc.version + 1}",
                document_type=original_doc.document_type,
                expiry_date=original_doc.expiry_date
            )
            
            # Set version information
            new_document.version = original_doc.version + 1
            new_document.parent_document_id = original_doc.id
            new_document.is_current_version = True
            
            self.db.session.commit()
            
            logger.info(f"New version created for document: {document_id}")
            return new_document
            
        except Exception as e:
            logger.error(f"Failed to create new version: {str(e)}")
            raise
    
    def search_documents(self, query, category=None, patient_id=None, date_from=None, date_to=None):
        """Search documents by content and metadata"""
        try:
            search_query = self.Document.query.filter(
                self.Document.status == 'active'
            )
            
            # Text search
            if query:
                search_query = search_query.filter(
                    self.Document.searchable_content.contains(query)
                )
            
            # Category filter
            if category:
                search_query = search_query.filter(self.Document.category == category)
            
            # Patient filter
            if patient_id:
                search_query = search_query.filter(self.Document.patient_id == patient_id)
            
            # Date range filter
            if date_from:
                search_query = search_query.filter(self.Document.created_at >= date_from)
            if date_to:
                search_query = search_query.filter(self.Document.created_at <= date_to)
            
            # Only current versions
            search_query = search_query.filter(self.Document.is_current_version == True)
            
            documents = search_query.order_by(self.Document.created_at.desc()).all()
            
            return documents
            
        except Exception as e:
            logger.error(f"Failed to search documents: {str(e)}")
            raise
    
    def check_expiring_documents(self, days_ahead=30):
        """Check for documents expiring soon"""
        try:
            expiry_date = datetime.now().date() + timedelta(days=days_ahead)
            
            expiring_docs = self.Document.query.filter(
                self.Document.expiry_date <= expiry_date,
                self.Document.expiry_date >= datetime.now().date(),
                self.Document.status == 'active'
            ).all()
            
            return expiring_docs
            
        except Exception as e:
            logger.error(f"Failed to check expiring documents: {str(e)}")
            raise
    
    def _generate_document_id(self):
        """Generate unique document ID"""
        timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
        random_suffix = hashlib.md5(os.urandom(16)).hexdigest()[:6]
        return f"DOC{timestamp}{random_suffix.upper()}"
    
    def _is_allowed_file(self, filename):
        """Check if file extension is allowed"""
        return '.' in filename and \
               filename.rsplit('.', 1)[1].lower() in self.config['allowed_extensions']
    
    def _calculate_file_hash(self, file_path):
        """Calculate SHA-256 hash of file"""
        hash_sha256 = hashlib.sha256()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(4096), b""):
                hash_sha256.update(chunk)
        return hash_sha256.hexdigest()
    
    def _scan_for_virus(self, file_path):
        """Scan file for viruses (placeholder - implement with actual antivirus)"""
        # This would integrate with ClamAV or similar
        # For now, return True (no virus found)
        return True
    
    def _encrypt_file(self, file_path):
        """Encrypt file using Fernet encryption"""
        try:
            with open(file_path, 'rb') as f:
                file_data = f.read()
            
            fernet = Fernet(self.encryption_key)
            encrypted_data = fernet.encrypt(file_data)
            
            encrypted_path = file_path.with_suffix(file_path.suffix + '.enc')
            with open(encrypted_path, 'wb') as f:
                f.write(encrypted_data)
            
            return encrypted_path
            
        except Exception as e:
            logger.error(f"Failed to encrypt file: {str(e)}")
            return None
    
    def _decrypt_file(self, encrypted_path):
        """Decrypt file"""
        try:
            with open(encrypted_path, 'rb') as f:
                encrypted_data = f.read()
            
            fernet = Fernet(self.encryption_key)
            decrypted_data = fernet.decrypt(encrypted_data)
            
            return decrypted_data
            
        except Exception as e:
            logger.error(f"Failed to decrypt file: {str(e)}")
            return None
    
    def _extract_text_ocr(self, file_path):
        """Extract text from image/PDF using OCR"""
        try:
            if file_path.suffix.lower() == '.pdf':
                # Extract text from PDF
                text = ""
                with open(file_path, 'rb') as f:
                    pdf_reader = PyPDF2.PdfReader(f)
                    for page in pdf_reader.pages:
                        text += page.extract_text()
                
                # If no text found, try OCR on PDF images
                if not text.strip():
                    # This would require pdf2image library
                    pass
                
                return text
            
            elif file_path.suffix.lower() in ['.jpg', '.jpeg', '.png', '.gif']:
                # OCR on image
                image = Image.open(file_path)
                text = pytesseract.image_to_string(image)
                return text
            
            return None
            
        except Exception as e:
            logger.error(f"Failed to extract text via OCR: {str(e)}")
            return None
    
    def _create_searchable_content(self, title, description=None, ocr_text=None):
        """Create searchable content string"""
        content_parts = [title or '']
        
        if description:
            content_parts.append(description)
        
        if ocr_text:
            content_parts.append(ocr_text)
        
        return ' '.join(content_parts).lower()
    
    def _render_template_content(self, template_content, data):
        """Render template with provided data"""
        try:
            from jinja2 import Template
            
            template = Template(template_content)
            return template.render(**data)
            
        except Exception as e:
            logger.error(f"Failed to render template: {str(e)}")
            return template_content
    
    def _get_or_create_encryption_key(self):
        """Get or create encryption key"""
        key_file = self.upload_folder / '.encryption_key'
        
        if key_file.exists():
            with open(key_file, 'rb') as f:
                return f.read()
        else:
            key = Fernet.generate_key()
            with open(key_file, 'wb') as f:
                f.write(key)
            return key
    
    def _get_client_ip(self):
        """Get client IP address"""
        if not request:
            return 'unknown'
        
        if request.environ.get('HTTP_X_FORWARDED_FOR') is None:
            return request.environ.get('REMOTE_ADDR', 'unknown')
        else:
            return request.environ['HTTP_X_FORWARDED_FOR']

def init_document_management(app, db):
    """Initialize document management system"""
    return DocumentManagement(db, app)